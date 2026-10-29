from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'manual' / 'HAFOEEP_Barista_Care_Manual_DE_Editable.docx'
INK='17191C'; INK2='25292E'; SILVER='C5C9CE'; MUTED='5D646C'; PAPER='F2F3F1'; WHITE='FFFFFF'; BORDER='D0D3D5'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}'); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')

def set_cell_border(cell, color=BORDER, size='6'):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:'+edge; element = borders.find(qn(tag))
        if element is None:
            element=OxmlElement(tag); borders.append(element)
        element.set(qn('w:val'),'single'); element.set(qn('w:sz'),size); element.set(qn('w:color'),color)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); tblHeader = OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'), 'true'); trPr.append(tblHeader)

def set_keep(p):
    pPr = p._p.get_or_add_pPr(); keep = OxmlElement('w:keepNext'); pPr.append(keep)

def set_run(run, size=None, bold=None, color=INK, font='Arial', italic=None):
    run.font.name = font
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None: run.italic = italic
    return run

def add_p(doc, text='', size=9, bold=False, color=INK, before=0, after=5, align=None, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    if align is not None: p.alignment = align
    if text:
        r=p.add_run(text); set_run(r,size=size,bold=bold,color=color)
    return p

def add_footer(section):
    footer=section.footer
    p=footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(0)
    r=p.add_run('HAFOEEP  ·  PRECISION COFFEE CARE'); set_run(r,size=6.5,bold=True,color=MUTED)

def add_cell_text(cell, lines, fill=None, margins=(120,140,120,140)):
    if fill: shade(cell,fill)
    set_cell_margins(cell,*margins)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text=''
    for idx, item in enumerate(lines):
        text, size, bold, color = item
        p=cell.paragraphs[0] if idx==0 else cell.add_paragraph()
        p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(2 if idx < len(lines)-1 else 0); p.paragraph_format.line_spacing=1.04
        r=p.add_run(text); set_run(r,size=size,bold=bold,color=color)

def add_card(doc, number, tag, title, body):
    tbl=doc.add_table(rows=1, cols=2); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER; tbl.autofit=False
    tbl.columns[0].width=Mm(26); tbl.columns[1].width=Mm(108)
    c0,c1=tbl.rows[0].cells; c0.width=Mm(26); c1.width=Mm(108)
    add_cell_text(c0,[(number,12,True,SILVER),(tag,6.5,True,WHITE)],fill=INK,margins=(90,80,90,80))
    add_cell_text(c1,[(title,9.3,True,INK),(body,7.5,False,INK)],fill=WHITE,margins=(100,140,100,140))
    for cell in (c0,c1): set_cell_border(cell)
    add_p(doc,'',size=2,after=1)

doc=Document()
sec=doc.sections[0]
sec.page_width=Mm(148); sec.page_height=Mm(210)
sec.top_margin=Mm(11); sec.bottom_margin=Mm(12); sec.left_margin=Mm(11); sec.right_margin=Mm(11)
sec.header_distance=Mm(4); sec.footer_distance=Mm(5)
add_footer(sec)

# PAGE 1 / COVER
p=add_p(doc,'HAFOEEP  /  PRECISION COFFEE CARE',size=8,bold=True,color=MUTED,after=7)
p.paragraph_format.keep_with_next=True
banner=doc.add_table(rows=1,cols=1); banner.alignment=WD_TABLE_ALIGNMENT.CENTER; banner.autofit=False; banner.columns[0].width=Mm(126)
cell=banner.cell(0,0); shade(cell,INK); set_cell_margins(cell,top=420,start=260,bottom=390,end=260); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell.text=''
p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(7)
r=p.add_run('H'); set_run(r,size=34,bold=True,color=SILVER)
p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(9)
r=p.add_run('HAFOEEP'); set_run(r,size=20,bold=True,color=WHITE)
p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(20)
r=p.add_run('PRECISION COFFEE CARE'); set_run(r,size=7,bold=True,color=SILVER)
p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(3)
r=p.add_run('BARISTA CARE'); set_run(r,size=23,bold=True,color=WHITE)
p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(8)
r=p.add_run('REINIGUNGSTUCH-SET'); set_run(r,size=15,bold=True,color=WHITE)
p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(0)
r=p.add_run('4 Tücher. 3 Aufgaben. Eine klare Routine.'); set_run(r,size=9,bold=False,color=SILVER)
add_p(doc,'REINIGEN  ·  TROCKNEN  ·  POLIEREN',size=9,bold=True,color=INK,before=12,after=7)
t=doc.add_table(rows=1,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
items=[('2 × CLEAN','Mikrofaser'),('1 × DRY','Waffelstruktur'),('1 × POLISH','Feines Poliertuch')]
for i,(a,b) in enumerate(items):
    t.columns[i].width=Mm(42); c=t.cell(0,i); add_cell_text(c,[(a,8,True,WHITE),(b,7,False,SILVER)],fill=INK,margins=(140,90,140,90)); set_cell_border(c,color=INK,size='4')
add_p(doc,'Für die tägliche Pflege deiner Siebträgermaschine und Kaffeemaschine.',size=8.5,color=MUTED,before=12,after=0)

# PAGE 2 / WHICH CLOTH
p=doc.add_paragraph(); p.paragraph_format.page_break_before=True; p.paragraph_format.space_after=Pt(7)
r=p.add_run('VIER TÜCHER. DREI AUFGABEN.'); set_run(r,size=16,bold=True,color=INK)
add_p(doc,'Wähle das Tuch nach dem Einsatzbereich. So bleiben Milch-, Kaffee- und Wasserreste besser voneinander getrennt.',size=8.2,color=MUTED,after=9)
add_card(doc,'01','CLEAN','Mikrofasertuch – Dampf & Milch','Ein Tuch für Milchreste an der Dampflanze. Hinweise des Maschinenherstellers beachten und heiße Bauteile vorsichtig handhaben.')
add_card(doc,'02','CLEAN','Mikrofasertuch – Kaffee','Das zweite Tuch für Siebträger, Kaffeerückstände und äußere Bereiche des Siebs. Möglichst getrennt vom Tuch für Milchreste halten.')
add_card(doc,'03','DRY','Waffelstruktur-Tuch','Zum Aufnehmen von überschüssigem Wasser an Tropfschale und Arbeitsfläche.')
add_card(doc,'04','POLISH','Feines Poliertuch','Für den letzten Wisch auf geeigneten glatten Oberflächen wie Edelstahl oder Glas. Empfindliche Flächen zuerst testen.')
add_p(doc,'Tipp: Die graue Farbe bleibt einheitlich – unterscheide die Tücher über ihre jeweilige Gewebestruktur und halte sie nach Einsatzbereich getrennt.',size=7.5,color=MUTED,before=5,after=0)

# PAGE 3 / ROUTINE
p=doc.add_paragraph(); p.paragraph_format.page_break_before=True; p.paragraph_format.space_after=Pt(5)
r=p.add_run('DIE BARISTA-ROUTINE'); set_run(r,size=16,bold=True,color=INK)
add_p(doc,'Nach jedem Kaffee – in vier einfachen Schritten.',size=8.2,color=MUTED,after=12)
steps=[
('01','REINIGEN','Milchreste und Kaffeerückstände mit den jeweils vorgesehenen Mikrofasertüchern entfernen.'),
('02','TROCKNEN','Überschüssiges Wasser mit dem Waffelstruktur-Tuch aufnehmen.'),
('03','POLIEREN','Geeignete glatte Außenflächen mit dem Poliertuch nachwischen.'),
('04','AUFHÄNGEN','Tücher nach Gebrauch ausspülen und gemäß Pflegeetikett waschen und trocknen lassen.'),
]
for n,h,b in steps:
    tbl=doc.add_table(rows=1,cols=2); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER; tbl.autofit=False; tbl.columns[0].width=Mm(16); tbl.columns[1].width=Mm(118)
    c0,c1=tbl.rows[0].cells; c0.width=Mm(16); c1.width=Mm(118)
    add_cell_text(c0,[(n,12,True,SILVER)],fill=INK,margins=(140,65,140,65))
    add_cell_text(c1,[(h,9.5,True,INK),(b,8,False,INK)],fill=PAPER,margins=(120,140,120,140))
    for c in (c0,c1): set_cell_border(c,color=BORDER,size='5')
    add_p(doc,'',size=2,after=4)
add_p(doc,'WICHTIG',size=9,bold=True,color=INK,before=7,after=3)
warning=doc.add_table(rows=1,cols=1); warning.alignment=WD_TABLE_ALIGNMENT.CENTER; cell=warning.cell(0,0)
add_cell_text(cell,[('Das Tuchset ersetzt keine vorgeschriebene Maschinenreinigung oder Entkalkung. Beachte immer die Pflege- und Sicherheitshinweise des Maschinenherstellers. Keine groben Partikel über empfindliche Oberflächen reiben.',8,False,INK)],fill='E2E4E3',margins=(150,170,150,170))
set_cell_border(cell,color='E2E4E3',size='4')

# PAGE 4 / CARE + QR PLACEHOLDER
p=doc.add_paragraph(); p.paragraph_format.page_break_before=True; p.paragraph_format.space_after=Pt(6)
r=p.add_run('PFLEGE & HINWEISE'); set_run(r,size=16,bold=True,color=INK)
care=[
('VOR DEM ERSTEN GEBRAUCH','Alle Tücher vor der ersten Anwendung waschen.'),
('RICHTIG WASCHEN','Pflegeetikett beachten. Keinen Weichspüler oder Bleichmittel verwenden.'),
('GETRENNT VERWENDEN','Die beiden Mikrofasertücher möglichst dauerhaft ihren Reinigungsbereichen zuordnen.'),
('OBERFLÄCHEN SCHONEN','Lose Partikel zuerst entfernen. Empfindliche oder beschichtete Flächen an unauffälliger Stelle testen.'),
]
for head,body in care:
    add_p(doc,head,size=8,bold=True,color=INK,before=4,after=2)
    add_p(doc,body,size=7.8,color=MUTED,after=5)
add_p(doc,'DIGITALE ANLEITUNG',size=9,bold=True,color=INK,before=6,after=3)
qr=doc.add_table(rows=1,cols=2); qr.alignment=WD_TABLE_ALIGNMENT.CENTER; qr.autofit=False; qr.columns[0].width=Mm(36); qr.columns[1].width=Mm(98)
c0,c1=qr.rows[0].cells; c0.width=Mm(36); c1.width=Mm(98)
add_cell_text(c0,[('QR',14,True,MUTED),('NACH DEPLOYMENT',5.8,True,MUTED)],fill='F7F7F6',margins=(260,80,260,80))
add_cell_text(c1,[('Anwendungsbeispiele und Tuch-Zuordnung',8.3,True,INK),('QR-Code erst nach Veröffentlichung der digitalen Anleitung einsetzen. Vor Druck den Scan mit einem echten Smartphone testen.',7.3,False,MUTED)],fill=PAPER,margins=(160,150,160,150))
for c in (c0,c1): set_cell_border(c,color=BORDER,size='5')
add_p(doc,'Lieferumfang: 4 Reinigungstücher mit Aufhängeschlaufen.',size=7.7,bold=True,color=INK,before=9,after=5)
add_p(doc,'Die Pflegehinweise am Produkt und die Bedienungsanleitung der Kaffeemaschine haben Vorrang.',size=7.2,color=MUTED,after=0)

# Global typography defaults, metadata
styles=doc.styles
styles['Normal'].font.name='Arial'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); styles['Normal'].font.size=Pt(8)
doc.core_properties.title='HAFOEEP Barista Care - German Instruction Manual'
doc.core_properties.subject='Four-piece barista cleaning towel set'
doc.core_properties.author='HAFOEEP'
doc.save(OUT)
print(OUT)
